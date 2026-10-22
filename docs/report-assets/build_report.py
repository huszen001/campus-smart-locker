# 把 实验报告.md 转成 实验报告.docx（只处理这份文稿用到的语法，不追求通用）
import re
import pathlib
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = pathlib.Path(__file__).resolve().parent.parent   # docs/
SRC = BASE / "实验报告.md"
OUT = BASE / "实验报告.docx"

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
MONO_FONT = "Consolas"
IMG_MAX_W = Cm(15.5)


def set_run_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), name)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(el, fill):
    pr = el.get_or_add_tcPr() if el.tag.endswith("tc") else el.get_or_add_pPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    pr.append(sh)


def add_rich(par, text, base_size=12, bold_all=False, font=BODY_FONT):
    """处理 **加粗** 与 `行内代码`"""
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            set_run_font(r, font, base_size, True)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            set_run_font(r, MONO_FONT, base_size - 1)
        else:
            r = par.add_run(part)
            set_run_font(r, font, base_size, bold_all)


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    sizes = {1: 16, 2: 14, 3: 12.5}
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(p, text, sizes[level], bold_all=True, font=HEAD_FONT)
    # 让标题进大纲（导航窗格 / 目录）
    p.style = doc.styles[f"Heading {min(level, 3)}"]
    for r in p.runs:
        set_run_font(r, HEAD_FONT, sizes[level], True, RGBColor(0, 0, 0))
    return p


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, cell in enumerate(header):
        c = t.cell(0, j)
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rich(p, cell, 10.5, bold_all=True, font=HEAD_FONT)
        shade(c._tc, "EFF6FF")
    for i, row in enumerate(body, start=1):
        for j, cell in enumerate(row):
            if j >= len(header):
                break
            c = t.cell(i, j)
            c.text = ""
            add_rich(c.paragraphs[0], cell, 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def add_code(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.6)
        shade(p._p, "F5F5F5")
        r = p.add_run(ln if ln else " ")
        set_run_font(r, MONO_FONT, 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc, path, caption, landscape=False, break_before=False):
    # 按图形的长宽比定宽度：先顶到版心宽度，太高就按高度反算，别让图缩成一条
    iw, ih = Image.open(path).size
    max_w = Cm(24.6) if landscape else IMG_MAX_W
    max_h = Cm(14.8) if landscape else Cm(19.0)
    w = max_w
    h = w * ih / iw
    if h > max_h:
        w = int(max_h * iw / ih)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    # 用段落属性翻页，而不是插一个分页符段落 —— 上一页排满时那个空段落会掉到新页，多出一张白纸
    p.paragraph_format.page_break_before = break_before
    p.add_run().add_picture(str(path), width=w)
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(10)
        add_rich(cp, caption, 10.5, bold_all=True, font=HEAD_FONT)


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def new_section(doc, landscape):
    from docx.enum.section import WD_SECTION, WD_ORIENT
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    if landscape:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21)
        s.top_margin = s.bottom_margin = Cm(2.0)
        s.left_margin = s.right_margin = Cm(2.2)
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(2.6)
    return s


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.6)
    add_page_number_footer(sec)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)

    lines = SRC.read_text(encoding="utf-8").split("\n")
    i = 0
    n = len(lines)
    cur_land = [False]   # 当前是否在横排节里

    def leave_landscape():
        # 横排节用完要切回竖排，否则后面的正文会排在横版纸上
        if cur_land[0]:
            new_section(doc, False)
            cur_land[0] = False

    while i < n:
        line = lines[i].rstrip()
        s = line.strip()

        if not s or s == "---":
            i += 1
            continue

        # 代码块
        if s.startswith("```"):
            leave_landscape()
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i].rstrip())
                i += 1
            i += 1
            add_code(doc, buf)
            continue

        # 图片（下一行通常是图题）
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", s)
        if m:
            land = m.group(1).strip().upper() == "L"
            img = (BASE / m.group(2)).resolve()
            cap = ""
            j = i + 1
            while j < n and not lines[j].strip():
                j += 1
            if j < n and lines[j].strip().startswith("**图"):
                cap = lines[j].strip().strip("*")
                i = j
            brk = False
            if land:
                if cur_land[0]:
                    brk = True             # 连着两张横排图：同一节里翻页，别来回切节
                else:
                    new_section(doc, True)
                    cur_land[0] = True
            else:
                leave_landscape()
            add_image(doc, img, cap, landscape=land, break_before=brk)
            i += 1
            continue

        # 标题
        if s.startswith("#"):
            leave_landscape()
            level = len(s) - len(s.lstrip("#"))
            text = s.lstrip("#").strip()
            if level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(6)
                add_rich(p, text, 20, bold_all=True, font=HEAD_FONT)
                i += 1
                continue
            add_heading(doc, text, level - 1)
            i += 1
            continue

        # 表格
        if s.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            leave_landscape()
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r"^[\s:|-]+$", "".join(cells)):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue

        # 列表
        if re.match(r"^[-*]\s+", s):
            leave_landscape()
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 1.5
            add_rich(p, re.sub(r"^[-*]\s+", "", s))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", s):
            leave_landscape()
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.line_spacing = 1.5
            add_rich(p, re.sub(r"^\d+\.\s+", "", s))
            i += 1
            continue

        # 正文
        leave_landscape()
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.85)
        add_rich(p, s)
        i += 1

    doc.save(OUT)
    print(f"已生成 {OUT}")


if __name__ == "__main__":
    main()
